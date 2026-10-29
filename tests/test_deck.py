import io

from core.deck import build_deck, ensure_template
from pptx import Presentation

DIGEST = {
    "title": "Engineering Digest",
    "period": "week of 27 Oct",
    "sections": [
        {"heading": "Highlights", "bullets": ["Shipped X", "Cut latency"]},
        {"heading": "Watch-outs", "bullets": ["Flaky queue"]},
    ],
}


def test_ensure_template_creates_16x9_file(tmp_path):
    path = ensure_template(str(tmp_path / "sub" / "t.pptx"))
    prs = Presentation(path)
    assert prs.slide_width > prs.slide_height  # widescreen


def test_deck_has_title_plus_one_slide_per_section(tmp_path):
    pptx = build_deck(DIGEST, template_path=ensure_template(str(tmp_path / "t.pptx")))
    prs = Presentation(io.BytesIO(pptx))
    assert len(prs.slides) == 1 + len(DIGEST["sections"])


def test_deck_title_slide_contains_digest_text(tmp_path):
    pptx = build_deck(DIGEST, template_path=ensure_template(str(tmp_path / "t.pptx")))
    prs = Presentation(io.BytesIO(pptx))
    texts = [shape.text_frame.text for shape in prs.slides[0].shapes if shape.has_text_frame]
    assert any("Engineering Digest" in t for t in texts)
